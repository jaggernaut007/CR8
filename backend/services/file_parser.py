"""File parsing and slide export for PDF and PPTX documents."""

import logging
import os
import shutil
import subprocess
import tempfile

import pymupdf

logger = logging.getLogger(__name__)

try:
    from langsmith import traceable
except ImportError:  # pragma: no cover
    def traceable(**kwargs):
        def decorator(func):
            return func
        return decorator


def extract_text(file_path: str) -> list[dict]:
    """Extract text from a PDF or PPTX file.

    Args:
        file_path: Absolute path to a ``.pdf`` or ``.pptx`` file.

    Returns:
        List of page dicts, each containing:
            - ``text`` -- extracted plain text for the page/slide.
            - ``source`` -- the original filename (basename only).
            - ``page`` -- 1-based page or slide number.

    Raises:
        ValueError: If the file extension is not ``.pdf`` or ``.pptx``.
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return _extract_pdf(file_path)
    if ext == ".pptx":
        return _extract_pptx(file_path)
    raise ValueError(f"Unsupported file type: {ext}")


def _extract_pdf(file_path: str) -> list[dict]:
    source = os.path.basename(file_path)
    pages = []
    with pymupdf.open(file_path) as doc:
        for i, page in enumerate(doc):
            text = page.get_text()
            pages.append({"text": text, "source": source, "page": i + 1})
    return pages


def _extract_pptx(file_path: str) -> list[dict]:
    from pptx import Presentation

    source = os.path.basename(file_path)
    prs = Presentation(file_path)
    pages = []
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        pages.append({"text": "\n".join(texts), "source": source, "page": i + 1})
    return pages


# ---------------------------------------------------------------------------
# Slide image export (for Kokoro video pipeline)
# ---------------------------------------------------------------------------


@traceable(run_type="tool", name="export_slides_as_images")
def export_slides_as_images(file_path: str, output_dir: str, dpi: int = 144) -> list[str]:
    """Export each page/slide of a PDF or PPTX as a PNG image.

    For PDF files, each page is rendered directly via PyMuPDF.
    For PPTX files, LibreOffice headless converts to PDF first, then
    each page is rendered.

    Args:
        file_path: Path to a ``.pdf`` or ``.pptx`` file.
        output_dir: Directory to write PNG files into.
        dpi: Render resolution (144 produces exactly 1920x1080 for
            standard 13.333" x 7.5" widescreen slides).

    Returns:
        Sorted list of PNG file paths (``slide_001.png``, ``slide_002.png``, ...).

    Raises:
        ValueError: If the file extension is unsupported or *output_dir*
            fails the path traversal check.
        RuntimeError: If LibreOffice conversion fails.
    """
    _validate_output_dir(output_dir)
    os.makedirs(output_dir, exist_ok=True)

    ext = os.path.splitext(file_path)[1].lower()
    logger.info("Exporting slides as images: %s (ext=%s, dpi=%d)", file_path, ext, dpi)
    if ext == ".pdf":
        paths = _export_pdf_pages(file_path, output_dir, dpi)
    elif ext == ".pptx":
        paths = _export_pptx_pages(file_path, output_dir, dpi)
    else:
        raise ValueError(f"Unsupported file type for image export: {ext}")
    logger.info("Exported %d slide images to %s", len(paths), output_dir)
    return paths


def _export_pdf_pages(file_path: str, output_dir: str, dpi: int = 144) -> list[str]:
    """Render each PDF page as a PNG via PyMuPDF."""
    paths: list[str] = []
    with pymupdf.open(file_path) as doc:
        for i, page in enumerate(doc):
            pix = page.get_pixmap(dpi=dpi)
            out_path = os.path.join(output_dir, f"slide_{i + 1:03d}.png")
            pix.save(out_path)
            paths.append(out_path)
    return paths


def _export_pptx_pages(file_path: str, output_dir: str, dpi: int = 144) -> list[str]:
    """Convert PPTX to PDF via LibreOffice headless, then render pages."""
    if not shutil.which("libreoffice"):
        raise RuntimeError(
            "libreoffice is not installed. PPTX slide export requires "
            "LibreOffice headless. On the CPU pipeline container, video "
            "generation (which needs slide images) should be offloaded to "
            "the GPU or CPU-video service."
        )
    with tempfile.TemporaryDirectory() as tmpdir:
        result = subprocess.run(
            [
                "libreoffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                tmpdir,
                os.path.abspath(file_path),
            ],
            capture_output=True,
            timeout=60,
            check=False,
        )
        if result.returncode != 0:
            raise RuntimeError(
                f"LibreOffice PPTX→PDF conversion failed (exit {result.returncode}): "
                f"{result.stderr.decode(errors='replace')}"
            )

        pdf_name = os.path.splitext(os.path.basename(file_path))[0] + ".pdf"
        pdf_path = os.path.join(tmpdir, pdf_name)
        if not os.path.exists(pdf_path):
            raise RuntimeError(
                f"LibreOffice did not produce expected PDF: {pdf_path}"
            )
        return _export_pdf_pages(pdf_path, output_dir, dpi)


def _validate_output_dir(output_dir: str) -> None:
    """Reject output directories that escape the working directory or /tmp."""
    resolved = os.path.realpath(output_dir)
    cwd = os.path.realpath(os.getcwd())
    tmp = os.path.realpath(tempfile.gettempdir())
    if not (
        resolved.startswith((cwd + os.sep, tmp + os.sep))
        or resolved in (cwd, tmp)
    ):
        raise ValueError(f"Path traversal detected: {output_dir}")
